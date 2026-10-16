import json
from datetime import time, timedelta
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from users.models import Profile
from .models import (
    Attendance, Book, Enrollment, Group, InboxMessage, LevelCalibration, Lesson, Payment, ScheduleSlot, Student,
)
from .views import get_current_period


def make_user(username, role=None, superuser=False):
    """Пользователь с ролью (role=None — без профиля)"""
    if superuser:
        user = User.objects.create_superuser(username, f'{username}@example.com', 'pass')
    else:
        user = User.objects.create_user(username, password='pass')
    if role:
        Profile.objects.create(user=user, role=role)
    return user


# Быстрый хешер паролей — тесты в разы быстрее
@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class CRMTestCase(TestCase):
    def setUp(self):
        self.admin = make_user('admin', 'admin')
        self.teacher = make_user('teacher', 'teacher')
        self.other_teacher = make_user('other', 'teacher')
        self.accountant = make_user('acc', 'accountant')
        self.manager = make_user('mgr', 'manager')

        self.group = Group.objects.create(name='Beginner', teacher=self.teacher)
        self.other_group = Group.objects.create(name='Альфа', teacher=self.other_teacher)
        self.student = Student.objects.create(name='Алия', price_per_lesson=Decimal('4000'))
        self.other_student = Student.objects.create(name='Тимур')
        self.enrollment = Enrollment.objects.create(student=self.student, group=self.group)
        self.other_enrollment = Enrollment.objects.create(student=self.other_student, group=self.other_group)

    def login(self, user):
        self.client.force_login(user)


class RoleTests(CRMTestCase):
    def test_user_without_profile_has_no_access(self):
        self.login(make_user('nobody'))
        for url in ['/dashboard/', '/students/add/', '/export/', f'/group/{self.group.id}/']:
            response = self.client.get(url)
            self.assertEqual(response.status_code, 403, url)
        response = self.client.post(f'/group/{self.group.id}/attendance/', {'student_id': self.student.id, 'status': 'present'})
        self.assertEqual(response.status_code, 403)

    def test_superuser_without_profile_is_admin(self):
        boss = make_user('boss', superuser=True)
        self.login(boss)
        self.assertEqual(self.client.get('/dashboard/').status_code, 200)
        self.assertEqual(self.client.get('/export/').status_code, 200)
        # Смена телефона создаёт профиль, но не понижает до учителя
        self.client.post('/profile/', {'action': 'change_phone', 'phone': '+7 700'})
        self.assertEqual(Profile.objects.get(user=boss).role, 'admin')

    def test_teacher_cannot_mark_other_group(self):
        self.login(self.teacher)
        response = self.client.post(f'/group/{self.other_group.id}/homework/', {'student_id': self.other_student.id, 'done': '1'})
        self.assertFalse(response.json()['success'])
        self.assertFalse(Attendance.objects.exists())

    def test_teacher_cannot_mark_student_from_another_group(self):
        self.login(self.teacher)
        response = self.client.post(f'/group/{self.group.id}/homework/', {'student_id': self.other_student.id, 'done': '1'})
        self.assertFalse(response.json()['success'])

    def test_attendance_status_is_validated(self):
        ScheduleSlot.objects.create(group=self.group, day_of_week=timezone.localdate().weekday(), start_time=time(10), end_time=time(11))
        self.login(self.teacher)
        response = self.client.post(f'/group/{self.group.id}/attendance/', {'student_id': self.student.id, 'status': 'hacked'})
        self.assertFalse(response.json()['success'])
        response = self.client.post(f'/group/{self.group.id}/attendance/', {'student_id': self.student.id, 'status': 'late'})
        self.assertTrue(response.json()['success'])
        self.assertEqual(Attendance.objects.get(student=self.student).status, 'late')

    def test_teacher_cannot_touch_books_and_calibration_of_other_group(self):
        self.login(self.teacher)
        self.client.post('/enrollment/toggle-book/', {'enrollment_ids': str(self.other_enrollment.id), 'action': 'need_book', 'comment': 'x'})
        self.other_enrollment.refresh_from_db()
        self.assertFalse(self.other_enrollment.book_needed)
        response = self.client.post('/calibration/create/', {'student_id': self.other_student.id, 'group_id': self.other_group.id, 'direction': 'up'})
        self.assertFalse(response.json()['success'])
        self.assertFalse(LevelCalibration.objects.exists())


class ManagerTests(CRMTestCase):
    def setUp(self):
        super().setUp()
        self.login(self.manager)

    def test_manager_sees_groups_students_and_payments(self):
        for url in ['/dashboard/', '/dashboard/?archived=1', '/students/', '/students/?archived=1', '/students/add/',
                    '/groups/add/', f'/group/{self.group.id}/', f'/group/{self.group.id}/edit/', f'/group/{self.other_group.id}/',
                    f'/student/{self.student.id}/edit/', f'/student/{self.student.id}/add-to-group/', '/individual/add/',
                    '/payments/', '/schedule/', '/export/']:
            self.assertEqual(self.client.get(url).status_code, 200, url)
        groups = self.client.get('/dashboard/').context['group_cards']
        self.assertEqual({g.name for g in groups}, {'Beginner', 'Альфа'})
        self.assertTrue(self.client.get(f'/group/{self.group.id}/').context['can_manage'])

    def test_manager_changes_students_and_marks_payment(self):
        self.client.post(f'/student/{self.student.id}/remove-from-group/{self.group.id}/')
        self.assertFalse(Enrollment.objects.filter(student=self.student, group=self.group).exists())
        self.client.post(f'/student/{self.other_student.id}/archive/')
        self.other_student.refresh_from_db()
        self.assertFalse(self.other_student.is_active)
        year, month = get_current_period()
        data = self.client.post(f'/payments/group/{self.other_group.id}/toggle/',
                                {'student_id': self.other_student.id, 'year': year, 'month': month}).json()
        self.assertTrue(data['success'])

    def test_role_accent_colors(self):
        violet, teal = '--accent: #8b5cf6;', '--accent: #0d9488;'
        response = self.client.get('/payments/')
        self.assertContains(response, teal)
        self.assertNotContains(response, violet)
        for admin in (self.admin, make_user('boss', superuser=True)):
            self.login(admin)
            response = self.client.get('/payments/')
            self.assertContains(response, violet)
            self.assertNotContains(response, teal)
        # Учитель и бухгалтер — оранжевый, без своей палитры
        for user in (self.teacher, self.accountant):
            self.login(user)
            response = self.client.get('/dashboard/')
            self.assertNotContains(response, violet)
            self.assertNotContains(response, teal)

    def test_manager_has_no_stats_broadcasts_books_and_rating(self):
        for url in ['/accountant-stats/', '/broadcast/', '/books-status/', '/star-rating/']:
            self.assertRedirects(self.client.get(url), '/dashboard/', fetch_redirect_response=False, msg_prefix=url)

    def test_create_manager_command(self):
        from django.core.management import call_command
        from django.core.management.base import CommandError
        call_command('create_manager', 'aksunkar', 'secret2026!')
        user = User.objects.get(username='aksunkar')
        self.assertEqual(user.profile.role, 'manager')
        self.assertTrue(self.client.login(username='aksunkar', password='secret2026!'))
        # Повторный запуск меняет пароль, а чужой логин не трогает
        call_command('create_manager', 'aksunkar', 'other2026!')
        self.assertTrue(self.client.login(username='aksunkar', password='other2026!'))
        with self.assertRaises(CommandError):
            call_command('create_manager', 'admin', 'secret2026!')
        self.assertEqual(Profile.objects.get(user=self.admin).role, 'admin')


class ArchiveTests(CRMTestCase):
    def setUp(self):
        super().setUp()
        Payment.objects.create(student=self.student, group=self.group, cycle_number=202610, amount=1000, is_paid=True)
        self.login(self.admin)

    def test_delete_links_do_not_work_by_get(self):
        for url in [f'/student/{self.student.id}/archive/', f'/group/{self.group.id}/archive/',
                    f'/student/{self.student.id}/remove-from-group/{self.group.id}/']:
            self.assertEqual(self.client.get(url).status_code, 405, url)
        self.student.refresh_from_db()
        self.assertTrue(self.student.is_active)

    def test_archive_student_keeps_history(self):
        self.client.post(f'/student/{self.student.id}/archive/')
        self.student.refresh_from_db()
        self.assertFalse(self.student.is_active)
        self.assertFalse(Enrollment.objects.filter(student=self.student).exists())
        self.assertTrue(Payment.objects.filter(student=self.student).exists())
        names = lambda url: [d['student'].name for d in self.client.get(url).context['student_data']]
        self.assertNotIn('Алия', names('/students/'))
        self.assertIn('Алия', names('/students/?archived=1'))

    def test_bulk_archive_students(self):
        self.assertEqual(self.client.get('/students/archive/').status_code, 405)
        self.client.post('/students/archive/', {'student_ids': [self.student.id, self.other_student.id]})
        for student in (self.student, self.other_student):
            student.refresh_from_db()
            self.assertFalse(student.is_active)
            self.assertIsNotNone(student.archived_at)
        self.assertFalse(Enrollment.objects.filter(student__in=[self.student, self.other_student]).exists())
        self.assertTrue(Payment.objects.filter(student=self.student).exists())

    def test_teacher_cannot_bulk_archive(self):
        self.login(self.teacher)
        self.client.post('/students/archive/', {'student_ids': [self.student.id]})
        self.student.refresh_from_db()
        self.assertTrue(self.student.is_active)

    def test_restore_student(self):
        self.client.post(f'/student/{self.student.id}/archive/')
        response = self.client.post(f'/student/{self.student.id}/restore/')
        self.assertRedirects(response, f'/student/{self.student.id}/add-to-group/')
        self.student.refresh_from_db()
        self.assertTrue(self.student.is_active)

    def test_adding_archived_student_to_group_restores_him(self):
        self.client.post(f'/student/{self.student.id}/archive/')
        self.client.post(f'/student/{self.student.id}/add-to-group/', {'groups': [self.group.id]})
        self.student.refresh_from_db()
        self.assertTrue(self.student.is_active)

    def test_archive_group_keeps_lessons_and_payments(self):
        Lesson.objects.create(group=self.group, date=timezone.localdate())
        self.client.post(f'/group/{self.group.id}/archive/')
        self.group.refresh_from_db()
        self.assertFalse(self.group.is_active)
        self.assertTrue(Lesson.objects.filter(group=self.group).exists())
        self.assertTrue(Payment.objects.filter(group=self.group).exists())
        names = lambda url: [g.name for g in self.client.get(url).context['group_cards']]
        self.assertNotIn('Beginner', names('/dashboard/'))
        self.assertIn('Beginner', names('/dashboard/?archived=1'))
        self.client.post(f'/group/{self.group.id}/restore/')
        self.group.refresh_from_db()
        self.assertTrue(self.group.is_active)

    def test_teacher_cannot_archive(self):
        self.login(self.teacher)
        self.client.post(f'/group/{self.group.id}/archive/')
        self.group.refresh_from_db()
        self.assertTrue(self.group.is_active)


class SearchTests(CRMTestCase):
    def test_special_characters_do_not_crash(self):
        self.login(self.admin)
        for url in ['/dashboard/?search=(', '/students/?search=[', '/payments/?year=2026&month=10&search=(*']:
            self.assertEqual(self.client.get(url).status_code, 200, url)

    def test_cyrillic_search_ignores_case(self):
        self.login(self.admin)
        self.assertContains(self.client.get('/dashboard/?search=АЛЬ'), 'Альфа')
        self.assertContains(self.client.get('/students/?search=алИ'), 'Алия')


class PaymentSummaryTests(CRMTestCase):
    def test_dashboard_counts_unpaid_students_without_payment_rows(self):
        third = Student.objects.create(name='Дана')
        Enrollment.objects.create(student=third, group=self.group)
        y, m = get_current_period(self.group.cycle_start_day)
        Payment.objects.create(student=self.student, group=self.group, cycle_number=y * 100 + m, amount=1, is_paid=True)
        self.login(self.admin)
        response = self.client.get('/dashboard/')
        # Оплатила одна из трёх (Алия; Дана и Тимур — нет)
        self.assertEqual(response.context['paid_count'], 1)
        self.assertEqual(response.context['unpaid_count'], 2)


class PaymentsPageTests(CRMTestCase):
    def setUp(self):
        super().setUp()
        Enrollment.objects.update(enrolled_at=timezone.now() - timedelta(days=400))
        ScheduleSlot.objects.create(group=self.group, day_of_week=0, start_time=time(10), end_time=time(11))
        self.year, self.month = get_current_period()
        self.cycle = self.year * 100 + self.month

    def page(self, **params):
        self.login(self.accountant)
        return self.client.get('/payments/', params)

    def row(self, response, name):
        return next(r for c in response.context['cards'] for r in c['rows'] if r['student'].name == name)

    def test_current_month_by_default(self):
        response = self.page()
        self.assertEqual((response.context['year'], response.context['month']), (self.year, self.month))
        self.assertEqual(len(response.context['year_bars']), 12)
        aliya = self.row(response, 'Алия')
        self.assertFalse(aliya['is_paid'])
        lessons = next(c['lessons'] for c in response.context['cards'] if c['group'] == self.group)
        self.assertEqual(aliya['amount'], Decimal('4000') * lessons)
        self.assertEqual(response.context['expected'], int(aliya['amount']))

    def test_paid_totals(self):
        Payment.objects.create(student=self.student, group=self.group, cycle_number=self.cycle,
                               amount=Decimal('16000'), paid_amount=Decimal('16000'), is_paid=True)
        response = self.page()
        self.assertEqual(response.context['collected'], 16000)
        self.assertEqual(response.context['paid_total'], 1)
        self.assertEqual(response.context['wait_total'], 1)
        selected = next(b for b in response.context['year_bars'] if b['is_selected'])
        self.assertEqual(selected['collected'], 16000)

    def test_student_who_left_keeps_paid_row(self):
        Payment.objects.create(student=self.student, group=self.group, cycle_number=self.cycle,
                               amount=Decimal('16000'), paid_amount=Decimal('16000'), is_paid=True)
        self.login(self.admin)
        self.client.post(f'/student/{self.student.id}/archive/')
        response = self.page()
        aliya = self.row(response, 'Алия')
        self.assertTrue(aliya['left'])
        self.assertEqual(response.context['collected'], 16000)

    def test_new_student_not_in_past_months(self):
        newbie = Student.objects.create(name='Дана', price_per_lesson=Decimal('3000'))
        Enrollment.objects.create(student=newbie, group=self.group)
        y, m = (self.year - 1, self.month)
        names = [r['student'].name for c in self.page(year=y, month=m).context['cards'] for r in c['rows']]
        self.assertIn('Алия', names)
        self.assertNotIn('Дана', names)

    def test_bad_params_fall_back_to_current_month(self):
        response = self.page(year='abc', month='13')
        self.assertEqual((response.context['year'], response.context['month']), (self.year, self.month))

    def test_teacher_has_no_access(self):
        self.login(self.teacher)
        self.assertRedirects(self.client.get('/payments/'), '/dashboard/', fetch_redirect_response=False)

    def test_toggle(self):
        self.login(self.accountant)
        url = f'/payments/group/{self.group.id}/toggle/'
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertEqual(self.client.post(url, {'student_id': self.student.id, 'year': 'x', 'month': 1}).status_code, 400)
        data = self.client.post(url, {'student_id': self.student.id, 'year': self.year, 'month': self.month}).json()
        self.assertTrue(data['is_paid'])
        self.assertIn('acc', data['paid_label'])
        payment = Payment.objects.get(student=self.student, group=self.group, cycle_number=self.cycle)
        self.assertEqual(data['amount'], int(payment.paid_amount))
        data = self.client.post(url, {'student_id': self.student.id, 'year': self.year, 'month': self.month}).json()
        self.assertFalse(data['is_paid'])
        self.assertEqual(data['paid_label'], '')


class GroupPageTests(CRMTestCase):
    def setUp(self):
        super().setUp()
        Enrollment.objects.update(enrolled_at=timezone.now() - timedelta(days=90))
        self.second = Student.objects.create(name='Бек')
        Enrollment.objects.filter(student=self.second).delete()
        Enrollment.objects.create(student=self.second, group=self.group, enrolled_at=timezone.now() - timedelta(days=90))
        today = timezone.localdate()
        # Три прошлых урока: Алия был, опоздал, не был; Бек — всегда был
        for days, status in [(9, 'present'), (6, 'late'), (3, 'absent')]:
            lesson = Lesson.objects.create(group=self.group, date=today - timedelta(days=days))
            Attendance.objects.create(lesson=lesson, student=self.student, status=status, homework_done=True, stars=2)
            Attendance.objects.create(lesson=lesson, student=self.second, status='present', stars=1)
        Lesson.objects.create(group=self.group, date=today - timedelta(days=2), is_cancelled=True)

    def data(self, response, name):
        return next(d for d in response.context['student_data'] if d['student'].name == name)

    def test_group_stats(self):
        self.login(self.teacher)
        response = self.client.get(f'/group/{self.group.id}/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['attendance_rate'], 83)   # 5 из 6 отметок
        self.assertEqual(response.context['homework_rate'], 50)
        self.assertEqual([r['rate'] for r in response.context['rhythm']], [100, 100, 50])
        aliya = self.data(response, 'Алия')
        self.assertEqual(aliya['dots'], ['present', 'late', 'absent'])
        self.assertEqual(aliya['rate'], 67)
        self.assertEqual(aliya['stars'], 6)
        self.assertEqual([d['student'].name for d in response.context['leaders']], ['Алия', 'Бек'])

    def test_teacher_sees_no_money(self):
        y, m = get_current_period(self.group.cycle_start_day)
        Payment.objects.create(student=self.student, group=self.group, cycle_number=y * 100 + m, amount=1, is_paid=True)
        self.login(self.teacher)
        response = self.client.get(f'/group/{self.group.id}/')
        self.assertFalse(response.context['can_money'])
        self.assertFalse(self.data(response, 'Алия')['is_paid'])
        self.login(self.accountant)
        response = self.client.get(f'/group/{self.group.id}/')
        self.assertTrue(self.data(response, 'Алия')['is_paid'])
        self.assertEqual((response.context['paid_count'], response.context['billable_count']), (1, 2))

    def test_lesson_today_is_created_on_lesson_day(self):
        today = timezone.localdate()
        ScheduleSlot.objects.create(group=self.group, day_of_week=today.weekday(), start_time=time(10), end_time=time(11))
        self.login(self.teacher)
        response = self.client.get(f'/group/{self.group.id}/')
        self.assertTrue(response.context['lesson_active'])
        self.assertEqual(response.context['today_slot'].start_time, time(10))
        self.assertEqual([d['attendance_status'] for d in response.context['student_data']], ['absent', 'absent'])
        self.assertContains(response, 'role="radiogroup"', count=2)

    def test_other_teacher_redirected(self):
        self.login(self.other_teacher)
        self.assertRedirects(self.client.get(f'/group/{self.group.id}/'), '/dashboard/', fetch_redirect_response=False)


class LessonTests(CRMTestCase):
    def setUp(self):
        super().setUp()
        Enrollment.objects.filter(id=self.enrollment.id).update(enrolled_at=timezone.now() - timedelta(days=60))
        self.lesson = Lesson.objects.create(group=self.group, date=timezone.localdate() - timedelta(days=3))
        Attendance.objects.create(lesson=self.lesson, student=self.student, status='late')

    def save(self, user, rows, topic=''):
        self.login(user)
        return self.client.post(f'/lesson/{self.lesson.id}/save/', json.dumps({'topic': topic, 'rows': rows}),
                                content_type='application/json')

    def test_history_counts_late_as_present(self):
        self.login(self.teacher)
        response = self.client.get(f'/group/{self.group.id}/history/')
        self.assertContains(response, 'Алия (опоздал)')
        self.assertEqual(response.context['lesson_data'][0]['present_count'], 1)

    def test_edit_past_lesson(self):
        response = self.save(self.teacher, [
            {'student_id': self.student.id, 'status': 'present', 'homework': True, 'stars': 500},
            {'student_id': self.other_student.id, 'status': 'present', 'homework': True, 'stars': 1},
            {'student_id': self.student.id, 'status': 'bogus'},
        ], topic='Past Simple')
        self.assertTrue(response.json()['success'])
        att = Attendance.objects.get(lesson=self.lesson, student=self.student)
        self.assertEqual((att.status, att.homework_done, att.stars), ('present', True, 100))
        # Чужой ученик не попал в урок
        self.assertFalse(Attendance.objects.filter(student=self.other_student).exists())
        self.lesson.refresh_from_db()
        self.assertEqual((self.lesson.topic, self.lesson.edited_by), ('Past Simple', self.teacher))

    def test_other_teacher_cannot_edit(self):
        response = self.save(self.other_teacher, [{'student_id': self.student.id, 'status': 'present'}])
        self.assertEqual(response.status_code, 403)
        self.assertEqual(Attendance.objects.get(student=self.student).status, 'late')

    def test_cancel_and_restore_lesson(self):
        self.login(self.teacher)
        self.client.post(f'/lesson/{self.lesson.id}/cancel/')
        self.lesson.refresh_from_db()
        self.assertTrue(self.lesson.is_cancelled)
        self.assertFalse(self.lesson.attendances.exists())
        self.assertFalse(self.save(self.teacher, [{'student_id': self.student.id, 'status': 'present'}]).json()['success'])

        self.client.post(f'/lesson/{self.lesson.id}/cancel/')
        self.lesson.refresh_from_db()
        self.assertFalse(self.lesson.is_cancelled)
        self.assertEqual(self.lesson.attendances.get().status, 'absent')

    def test_cancelled_today_lesson_blocks_marks(self):
        today = Lesson.objects.create(group=self.group, date=timezone.localdate(), is_cancelled=True)
        ScheduleSlot.objects.create(group=self.group, day_of_week=today.date.weekday(), start_time=time(10), end_time=time(11))
        self.login(self.teacher)
        for url, data in [('attendance', {'student_id': self.student.id, 'status': 'present'}),
                          ('homework', {'student_id': self.student.id, 'done': '1'}),
                          ('stars', {'stars': f'{self.student.id}:2'})]:
            self.assertFalse(self.client.post(f'/group/{self.group.id}/{url}/', data).json()['success'], url)

    def test_add_past_lesson(self):
        self.login(self.teacher)
        future = timezone.localdate() + timedelta(days=1)
        self.client.post(f'/group/{self.group.id}/history/add/', {'date': future.isoformat()})
        self.assertFalse(Lesson.objects.filter(date=future).exists())

        day = timezone.localdate() - timedelta(days=10)
        response = self.client.post(f'/group/{self.group.id}/history/add/', {'date': day.isoformat()})
        lesson = Lesson.objects.get(group=self.group, date=day)
        self.assertIn(f'?open={lesson.id}', response['Location'])
        self.assertEqual(lesson.attendances.get().student, self.student)

    def test_lesson_day_comes_from_schedule_slots(self):
        ScheduleSlot.objects.create(group=self.group, day_of_week=timezone.localdate().weekday(), start_time=time(10), end_time=time(11))
        self.login(self.teacher)
        self.client.get(f'/group/{self.group.id}/')
        self.assertTrue(Lesson.objects.filter(group=self.group, date=timezone.localdate()).exists())

    def test_archived_group_does_not_create_lessons(self):
        ScheduleSlot.objects.create(group=self.group, day_of_week=timezone.localdate().weekday(), start_time=time(10), end_time=time(11))
        self.group.is_active = False
        self.group.save()
        self.login(self.admin)
        self.client.get(f'/group/{self.group.id}/')
        self.assertFalse(Lesson.objects.filter(group=self.group, date=timezone.localdate()).exists())


class CleanupTests(CRMTestCase):
    def test_old_payment_pages_are_gone(self):
        self.login(self.admin)
        self.assertEqual(self.client.get('/payments-list/').status_code, 404)
        self.assertRedirects(self.client.get('/payments-management/'), '/payments/')

    def test_logout_drops_pending_messages(self):
        self.login(self.teacher)
        self.client.get(f'/group/{self.other_group.id}/')  # «нет доступа» — сообщение
        self.client.post('/logout/')
        self.assertFalse(list(self.client.get('/').context['messages']))


class IndividualTests(CRMTestCase):
    def setUp(self):
        super().setUp()
        self.solo = Group.objects.create(name='Индивид. Алия', group_type='individual', teacher=self.teacher)
        ScheduleSlot.objects.create(group=self.solo, day_of_week=0, start_time=time(16), end_time=time(17))
        Enrollment.objects.create(student=self.student, group=self.solo)
        self.other_solo = Group.objects.create(name='Индивид. Тимур', group_type='individual', teacher=self.other_teacher)
        Enrollment.objects.create(student=self.other_student, group=self.other_solo)

    def test_tab_lists_individual_students(self):
        self.login(self.admin)
        response = self.client.get('/individual/')
        self.assertEqual(response.status_code, 200)
        names = sorted(c['student'].name for c in response.context['cards'])
        self.assertEqual(names, ['Алия', 'Тимур'])
        self.assertTrue(response.context['can_money'])

    def test_teacher_sees_only_own_students_without_money(self):
        self.login(self.teacher)
        response = self.client.get('/individual/')
        self.assertEqual([c['student'].name for c in response.context['cards']], ['Алия'])
        self.assertFalse(response.context['can_money'])
        self.assertNotContains(response, '/payments/group/')

    def test_dashboard_hides_individual_groups(self):
        self.login(self.admin)
        response = self.client.get('/dashboard/')
        shown = {g.id for g in response.context['group_cards']}
        self.assertIn(self.group.id, shown)
        self.assertNotIn(self.solo.id, shown)
        self.assertEqual(response.context['individual_count'], 2)

    def test_add_individual(self):
        newbie = Student.objects.create(name='Дана')
        self.login(self.admin)
        response = self.client.post('/individual/add/', {
            'student': newbie.id, 'teacher': self.teacher.id, 'price': '5 000',
            'day_1': '1', 'start_1': '18:00', 'end_1': '19:00',
            'day_3': '1', 'start_3': '18:00', 'end_3': 'oops',
            'cycle_start_day': '5',
        })
        self.assertRedirects(response, '/individual/')
        group = Group.objects.get(enrollments__student=newbie)
        self.assertEqual(group.group_type, 'individual')
        self.assertEqual(group.teacher, self.teacher)
        self.assertEqual(group.cycle_start_day, 5)
        self.assertEqual(list(group.schedule_slots.values_list('day_of_week', flat=True)), [1])
        newbie.refresh_from_db()
        self.assertEqual(newbie.price_per_lesson, Decimal('5000'))

    def test_add_individual_admin_only(self):
        newbie = Student.objects.create(name='Дана')
        self.login(self.teacher)
        self.client.post('/individual/add/', {'student': newbie.id, 'teacher': self.teacher.id})
        self.assertFalse(Enrollment.objects.filter(student=newbie).exists())

    def test_split_shared_group_moves_history(self):
        shared = Group.objects.create(name='individual', teacher=self.teacher, cycle_start_day=5)
        ScheduleSlot.objects.create(group=shared, day_of_week=2, start_time=time(15), end_time=time(16))
        Enrollment.objects.create(student=self.student, group=shared)
        Enrollment.objects.create(student=self.other_student, group=shared)
        day = timezone.localdate() - timedelta(days=7)
        # Два урока в один день не должны склеиться в один
        first = Lesson.objects.create(group=shared, date=day, topic='Present Simple')
        second = Lesson.objects.create(group=shared, date=day, topic='Доп. урок')
        for lesson in (first, second):
            Attendance.objects.create(lesson=lesson, student=self.student, status='present', stars=2)
        Attendance.objects.create(lesson=first, student=self.other_student, status='absent')
        Payment.objects.create(student=self.student, group=shared, cycle_number=202610, amount=Decimal('8000'), is_paid=True)
        LevelCalibration.objects.create(student=self.student, group=shared, direction='up')

        self.login(self.admin)
        self.assertEqual(self.client.get('/individual/').context['shared_groups'], [shared])
        response = self.client.post(f'/individual/split/{shared.id}/')
        self.assertRedirects(response, '/individual/')

        shared.refresh_from_db()
        self.assertFalse(shared.is_active)
        self.assertFalse(shared.enrollments.exists())
        aliya = Group.objects.get(enrollments__student=self.student, group_type='individual', is_active=True,
                                  schedule_slots__day_of_week=2)
        self.assertEqual(aliya.teacher, self.teacher)
        self.assertEqual(aliya.cycle_start_day, 5)
        self.assertEqual(sorted(aliya.lessons.values_list('topic', flat=True)), ['Present Simple', 'Доп. урок'])
        self.assertEqual(Attendance.objects.filter(student=self.student, lesson__group=aliya).count(), 2)
        self.assertTrue(Payment.objects.filter(student=self.student, group=aliya, is_paid=True).exists())
        self.assertTrue(LevelCalibration.objects.filter(student=self.student, group=aliya).exists())
        timur = Group.objects.get(enrollments__student=self.other_student, schedule_slots__day_of_week=2)
        self.assertEqual(list(timur.lessons.values_list('topic', flat=True)), ['Present Simple'])
        self.assertFalse(Attendance.objects.filter(lesson__group=shared).exists())

    def test_split_requires_post_and_admin(self):
        shared = Group.objects.create(name='individual', teacher=self.teacher)
        Enrollment.objects.create(student=self.student, group=shared)
        self.login(self.admin)
        self.assertEqual(self.client.get(f'/individual/split/{shared.id}/').status_code, 405)
        self.login(self.teacher)
        self.client.post(f'/individual/split/{shared.id}/')
        shared.refresh_from_db()
        self.assertTrue(shared.is_active)

    def test_single_student_group_just_moves(self):
        lone = Group.objects.create(name='Индивид. Марат', teacher=self.teacher)
        Enrollment.objects.create(student=self.student, group=lone)
        lesson = Lesson.objects.create(group=lone, date=timezone.localdate())
        self.login(self.admin)
        self.client.post(f'/individual/split/{lone.id}/')
        lone.refresh_from_db()
        self.assertEqual(lone.group_type, 'individual')
        self.assertTrue(lone.is_active)
        self.assertTrue(Lesson.objects.filter(id=lesson.id, group=lone).exists())


class EnglishTests(CRMTestCase):
    """Английская версия: страница приходит с сервера уже переведённой (dashboard/i18n.py)"""

    def test_phrases(self):
        from . import i18n
        self.assertEqual(i18n.translate_text('  Сохранить '), '  Save ')
        self.assertEqual(i18n.translate_text('Собрано за октябрь'), 'Collected in October')
        self.assertEqual(i18n.translate_text('Следующий урок в пт 10:00'), 'Next lesson on Fri 10:00')
        self.assertEqual(i18n.translate_text('1 учеников'), '1 student')
        self.assertEqual(i18n.translate_text('5 учеников'), '5 students')
        self.assertEqual(i18n.translate_text('Октябрь 2026'), 'October 2026')
        self.assertEqual(i18n.translate_text('Оплата: Алия'), 'Payment: Алия')
        self.assertEqual(i18n.translate_text('Алия'), 'Алия')

    def test_html(self):
        from . import i18n
        page = ('<html lang="ru"><title>Оплаты — Enjoy English CRM</title>'
                '<p class="notranslate">Сохранить</p><input placeholder="Найти ученика…">'
                '<textarea>Сохранить</textarea>'
                "<script>alert('У вас нет доступа'); const k = 'л';</script>"
                '<button onclick="confirm(\'Выберите направление\')">Сохранить</button></html>')
        html = i18n.translate_html(page)
        self.assertIn('<html lang="en">', html)
        self.assertIn('<title>Payments — Enjoy English CRM</title>', html)
        self.assertIn('<p class="notranslate">Сохранить</p>', html)
        self.assertIn('placeholder="Find a student…"', html)
        self.assertIn('<textarea>Сохранить</textarea>', html)
        self.assertIn("alert('You don\\'t have access')", html)
        self.assertIn("const k = 'л'", html)
        self.assertIn('Choose a direction', html)
        self.assertIn('>Save</button>', html)

    def test_page_comes_translated(self):
        self.login(self.admin)
        self.assertContains(self.client.get('/dashboard/'), 'Посещаемость')
        self.client.cookies['ee_lang'] = 'en'
        response = self.client.get('/dashboard/')
        self.assertContains(response, '<html lang="en"')
        self.assertContains(response, 'Attendance')
        self.assertNotContains(response, 'Посещаемость')
        self.assertNotContains(response, 'translate.google.com')

    def test_ajax_errors_translated(self):
        self.login(self.teacher)
        self.client.cookies['ee_lang'] = 'en'
        data = self.client.post(f'/group/{self.group.id}/attendance/', {'student_id': self.student.id, 'status': 'bogus'}).json()
        self.assertIn(data['error'], ['There is no lesson today', 'Invalid status'])

    def test_english_chosen_in_google_is_kept(self):
        self.login(self.admin)
        self.client.cookies['googtrans'] = '/ru/en'
        self.assertContains(self.client.get('/dashboard/'), '<html lang="en"')
        self.client.cookies['ee_lang'] = 'ru'
        self.assertContains(self.client.get('/dashboard/'), '<html lang="ru"')

    def test_kazakh_still_uses_google(self):
        self.client.cookies['googtrans'] = '/ru/kk'
        self.assertContains(self.client.get('/'), 'translate.google.com')
        self.client.cookies['googtrans'] = ''
        self.assertNotContains(self.client.get('/'), 'translate.google.com')

    def test_every_phrase_has_translation(self):
        """Новый русский текст в шаблонах — добавьте перевод в dashboard/locale_en.py"""
        from . import i18n
        today = timezone.localdate()
        ScheduleSlot.objects.create(group=self.group, day_of_week=today.weekday(), start_time=time(10), end_time=time(11))
        names = {'Алия', 'Тимур', 'Альфа', 'алия', 'тимур', 'альфа'}
        i18n.missing = set()
        try:
            self.client.cookies['ee_lang'] = 'en'
            self.client.get('/')
            for user in [self.admin, self.teacher, self.accountant, self.manager]:
                self.login(user)
                for url in ['/dashboard/', '/students/', '/students/add/', '/schedule/', '/individual/', '/payments/',
                            '/accountant-stats/', '/groups/add/', '/books-status/', '/star-rating/', '/profile/',
                            '/broadcast/', '/broadcast/?kind=news', '/broadcast/?kind=payment&audience=teacher',
                            f'/group/{self.group.id}/', f'/group/{self.group.id}/history/', f'/group/{self.group.id}/edit/',
                            f'/students/{self.student.id}/', f'/student/{self.student.id}/edit/']:
                    self.client.get(url)
            # Одна буква — инициалы на аватарах и сочетание Ctrl+Л
            missing = {phrase for phrase in i18n.missing - names if len(phrase) > 1}
        finally:
            i18n.missing = None
        self.assertFalse(missing, f'Нет английского перевода (добавьте в dashboard/locale_en.py): {sorted(missing)}')


class BroadcastTests(CRMTestCase):
    """Рассылки в WhatsApp: суммы, получатели, отправка"""

    def setUp(self):
        super().setUp()
        Enrollment.objects.update(enrolled_at=timezone.now() - timedelta(days=90))
        for g in (self.group, self.other_group):
            ScheduleSlot.objects.create(group=g, day_of_week=0, start_time=time(10), end_time=time(11))
        self.student.parent_name = 'Айгуль Маратовна'
        self.student.parent_phone = '8 (700) 111-22-33'
        self.student.save()
        self.other_student.phone = '+7 701 444 55 66'
        self.other_student.price_per_lesson = Decimal('3000')
        self.other_student.save()
        from .broadcasts import default_month
        self.year, self.month = default_month()

    def recipients(self, kind='payment', audience='parent', **kw):
        from .broadcasts import build_recipients
        return build_recipients(kind, audience, kw.get('groups'), kw.get('type', 'all'), kw.get('unpaid', True),
                                self.year, self.month, kw.get('mode', 'manual'))

    def lessons(self, group):
        from .views import count_lessons_in_period, get_period
        return count_lessons_in_period(group, *get_period(self.year, self.month, group.cycle_start_day))

    def test_normalize_phone(self):
        from .whatsapp import normalize_phone
        self.assertEqual(normalize_phone('8 (700) 111-22-33'), '77001112233')
        self.assertEqual(normalize_phone('+7 700 111 22 33'), '77001112233')
        self.assertEqual(normalize_phone('7001112233'), '77001112233')
        self.assertIsNone(normalize_phone('12-34'))
        self.assertIsNone(normalize_phone(''))

    def test_payment_amounts_for_parents(self):
        rows = {r['name']: r for r in self.recipients()}
        aliya = rows['Айгуль Маратовна']
        expected = Decimal('4000') * self.lessons(self.group)
        self.assertEqual(aliya['phone'], '77001112233')
        self.assertEqual(Decimal(str(aliya['amount'])), expected)
        self.assertIn('Здравствуйте, Айгуль!', aliya['vars']['приветствие'])
        self.assertIn('Алия — Beginner', aliya['vars']['детали'])
        # У Тимура нет номера родителя — отправим на номер ученика
        timur = rows['Родитель: Тимур']
        self.assertTrue(timur['ok'])
        self.assertIn('номер ученика', timur['warn'])

    def test_paid_students_are_skipped(self):
        cycle = self.year * 100 + self.month
        Payment.objects.create(student=self.student, group=self.group, cycle_number=cycle, amount=1, is_paid=True)
        names = [r['name'] for r in self.recipients()]
        self.assertNotIn('Айгуль Маратовна', names)
        names = [r['name'] for r in self.recipients(unpaid=False)]
        self.assertIn('Айгуль Маратовна', names)

    def test_siblings_share_one_message(self):
        sister = Student.objects.create(name='Дана', parent_phone='+77001112233', price_per_lesson=Decimal('2000'))
        Enrollment.objects.create(student=sister, group=self.other_group, enrolled_at=timezone.now() - timedelta(days=90))
        rows = [r for r in self.recipients() if r['phone'] == '77001112233']
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['vars']['имя'], 'Алия и Дана')
        total = Decimal('4000') * self.lessons(self.group) + Decimal('2000') * self.lessons(self.other_group)
        self.assertEqual(Decimal(str(rows[0]['amount'])), total)

    def test_no_phone_or_price_cannot_be_sent(self):
        self.other_student.phone = ''
        self.other_student.save()
        timur = next(r for r in self.recipients() if 'Тимур' in r['name'])
        self.assertFalse(timur['ok'])
        self.student.price_per_lesson = 0
        self.student.save()
        aliya = next(r for r in self.recipients() if r['name'] == 'Айгуль Маратовна')
        self.assertFalse(aliya['ok'])

    def test_teacher_gets_list_of_unpaid(self):
        Profile.objects.filter(user=self.teacher).update(phone='+7 702 000 00 01')
        rows = {r['name']: r for r in self.recipients(audience='teacher')}
        self.assertEqual(rows['teacher']['phone'], '77020000001')
        self.assertIn('Алия', rows['teacher']['vars']['детали'])
        self.assertNotIn('Тимур', rows['teacher']['vars']['детали'])
        # У второго учителя нет телефона — сообщение ему не уйдёт
        self.assertFalse(rows['other']['ok'])

    def test_group_chats(self):
        self.group.whatsapp_chat_id = '120363000000000001@g.us'
        self.group.save()
        solo = Group.objects.create(name='Индивид. Алия', group_type='individual', teacher=self.teacher)
        rows = {r['name']: r for r in self.recipients(kind='news', audience='group')}
        self.assertEqual(rows['Beginner']['chat_id'], '120363000000000001@g.us')
        self.assertIn('скопировать', rows['Альфа']['warn'])
        self.assertNotIn(solo.name, rows)

    def test_render_text(self):
        from .broadcasts import render_text
        self.assertEqual(render_text('{приветствие}\n\n\n{сумма} {неизвестно}', {'приветствие': 'Привет!', 'сумма': '5 ₸'}),
                         'Привет!\n\n5 ₸ {неизвестно}')

    def create(self, user, **data):
        self.login(user)
        body = {'kind': 'payment', 'audience': 'parent', 'year': self.year, 'month': self.month,
                'text': '{приветствие} К оплате {сумма} за {месяц}.'}
        body.update(data)
        return self.client.post('/broadcast/create/', json.dumps(body), content_type='application/json')

    def test_create_broadcast(self):
        from .models import Broadcast
        response = self.create(self.accountant, excluded=['p77014445566'])
        self.assertTrue(response.json()['success'])
        broadcast = Broadcast.objects.get()
        self.assertEqual(broadcast.mode, 'manual')
        message = broadcast.messages.get()
        self.assertEqual(message.phone, '77001112233')
        self.assertTrue(message.text.startswith('Здравствуйте, Айгуль! К оплате '))
        self.assertEqual(self.client.get(f'/broadcast/{broadcast.id}/').status_code, 200)

    def test_connected_whatsapp_starts_sending_at_once(self):
        from unittest import mock
        from .models import WhatsAppSettings
        self.assertNotIn('start=1', self.create(self.admin).json()['url'])
        WhatsAppSettings.objects.update_or_create(pk=1, defaults={'provider': 'green', 'instance_id': '1101', 'token': 'secret'})
        with mock.patch('dashboard.whatsapp.GreenAPI.state', return_value='authorized'):
            data = self.create(self.admin).json()
        self.assertTrue(data['url'].endswith(f'/broadcast/{data["id"]}/?start=1'))

    def test_news_needs_text_and_teacher_has_no_access(self):
        from .broadcasts import TEMPLATES
        response = self.create(self.admin, kind='news', audience='group', text=TEMPLATES[('news', 'group')])
        self.assertFalse(response.json()['success'])
        self.assertEqual(self.create(self.teacher).status_code, 403)
        self.login(self.teacher)
        self.assertRedirects(self.client.get('/broadcast/'), '/dashboard/', fetch_redirect_response=False)

    def test_send_through_green_api(self):
        from unittest import mock
        from .models import BroadcastMessage, WhatsAppSettings
        from .whatsapp import WhatsAppError
        WhatsAppSettings.objects.create(pk=1, provider='green', instance_id='1101', token='secret')
        self.create(self.admin, excluded=['p77014445566'])
        message = BroadcastMessage.objects.get()
        with mock.patch('dashboard.whatsapp.GreenAPI.send', return_value='BAE5') as send:
            data = self.client.post(f'/broadcast/message/{message.id}/send/').json()
            self.assertEqual(data['status'], 'sent')
            send.assert_called_once_with('77001112233@c.us', message.text)
            # Повторно не отправляется
            self.client.post(f'/broadcast/message/{message.id}/send/')
            self.assertEqual(send.call_count, 1)
        message.status = 'pending'
        message.save()
        with mock.patch('dashboard.whatsapp.GreenAPI.send', side_effect=WhatsAppError('Нет связи с Green API')):
            data = self.client.post(f'/broadcast/message/{message.id}/send/').json()
        self.assertEqual(data['status'], 'failed')
        message.refresh_from_db()
        self.assertEqual(message.error, 'Нет связи с Green API')

    def test_manual_mark(self):
        from .models import BroadcastMessage
        self.create(self.admin)
        message = BroadcastMessage.objects.first()
        self.assertEqual(self.client.post(f'/broadcast/message/{message.id}/mark/', {'status': 'sent'}).json()['status'], 'sent')
        message.refresh_from_db()
        self.assertEqual(message.status, 'sent')
        self.assertEqual(self.client.post(f'/broadcast/message/{message.id}/mark/', {'status': 'hack'}).status_code, 400)

    def test_bind_group_chat(self):
        self.login(self.admin)
        url = '/broadcast/whatsapp/bind/'
        self.assertEqual(self.client.post(url, {'group_id': self.group.id, 'chat_id': '77001112233@c.us'}).status_code, 400)
        self.assertTrue(self.client.post(url, {'group_id': self.group.id, 'chat_id': '120363000000000001@g.us'}).json()['success'])
        self.group.refresh_from_db()
        self.assertEqual(self.group.whatsapp_chat_id, '120363000000000001@g.us')
        self.login(self.accountant)
        self.assertEqual(self.client.post(url, {'group_id': self.group.id, 'chat_id': ''}).status_code, 403)

    def test_green_api_client(self):
        import io
        from unittest import mock
        from .models import WhatsAppSettings
        from .whatsapp import GreenAPI

        class Response(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

        settings = WhatsAppSettings(provider='green', api_url='https://7103.api.greenapi.com/', instance_id='1101', token='tok')
        with mock.patch('urllib.request.urlopen', return_value=Response(b'{"idMessage": "BAE5"}')) as urlopen:
            self.assertEqual(GreenAPI(settings).send('77001112233@c.us', 'Привет'), 'BAE5')
        request = urlopen.call_args[0][0]
        self.assertEqual(request.full_url, 'https://7103.api.greenapi.com/waInstance1101/sendMessage/tok')
        self.assertEqual(json.loads(request.data), {'chatId': '77001112233@c.us', 'message': 'Привет'})
        contacts = b'[{"id": "1@g.us", "name": "B", "type": "group"}, {"id": "2@c.us", "name": "A", "type": "user"}]'
        with mock.patch('urllib.request.urlopen', return_value=Response(contacts)):
            self.assertEqual(GreenAPI(settings).groups(), [('1@g.us', 'B')])

    def test_waha_client(self):
        import io
        import urllib.error
        from unittest import mock
        from .models import WhatsAppSettings
        from .whatsapp import WAHA

        class Response(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

        settings = WhatsAppSettings(provider='waha', api_url='http://localhost:3000/', token='key')
        self.assertTrue(settings.is_configured)
        with mock.patch('urllib.request.urlopen', return_value=Response(b'{"id": "true_77001112233@c.us_3EB0"}')) as urlopen:
            self.assertEqual(WAHA(settings).send('77001112233@c.us', 'Привет'), 'true_77001112233@c.us_3EB0')
        request = urlopen.call_args[0][0]
        self.assertEqual(request.full_url, 'http://localhost:3000/api/sendText')
        self.assertEqual(request.get_header('X-api-key'), 'key')
        self.assertEqual(json.loads(request.data), {'session': 'default', 'chatId': '77001112233@c.us', 'text': 'Привет'})
        # Группы: у разных движков WAHA разный формат ответа
        webjs = b'[{"id": {"_serialized": "1@g.us"}, "name": "B"}, {"id": "2@c.us", "name": "A"}]'
        noweb = b'{"3@g.us": {"id": "3@g.us", "subject": "C"}}'
        with mock.patch('urllib.request.urlopen', side_effect=[Response(webjs), Response(noweb)]):
            self.assertEqual(WAHA(settings).groups(), [('1@g.us', 'B')])
            self.assertEqual(WAHA(settings).groups(), [('3@g.us', 'C')])
        # Сессии ещё нет — WAHA отвечает 404
        not_found = urllib.error.HTTPError('http://localhost:3000/api/sessions/default', 404, 'Not Found', {}, None)
        with mock.patch('urllib.request.urlopen', side_effect=not_found):
            self.assertEqual(WAHA(settings).state(), 'missing')

    def test_connect_waha_and_send(self):
        from unittest import mock
        from .models import BroadcastMessage, WhatsAppSettings
        self.login(self.admin)
        with mock.patch('dashboard.whatsapp.WAHA.state', return_value='missing'), \
                mock.patch('dashboard.whatsapp.WAHA.start') as start:
            response = self.client.post('/broadcast/whatsapp/', {'provider': 'waha', 'api_url': 'http://localhost:3000', 'token': 'key'})
        self.assertRedirects(response, '/broadcast/?setup=1', fetch_redirect_response=False)
        start.assert_called_once()
        settings = WhatsAppSettings.load()
        self.assertEqual((settings.provider, settings.token), ('waha', 'key'))
        # Экран подключения показывает QR-код
        with mock.patch('dashboard.whatsapp.WAHA.state', return_value='SCAN_QR_CODE'), \
                mock.patch('dashboard.whatsapp.WAHA.qr', return_value='data:image/png;base64,AAA'):
            data = self.client.get('/broadcast/whatsapp/qr/').json()
        self.assertEqual(data, {'state': 'SCAN_QR_CODE', 'qr': 'data:image/png;base64,AAA'})
        # QR отсканирован — рассылка уходит сама через WAHA
        with mock.patch('dashboard.whatsapp.WAHA.state', return_value='WORKING'):
            self.assertEqual(self.client.get('/broadcast/whatsapp/qr/').json()['state'], 'WORKING')
            created = self.create(self.admin, excluded=['p77014445566']).json()
        self.assertIn('?start=1', created['url'])
        message = BroadcastMessage.objects.get()
        with mock.patch('dashboard.whatsapp.WAHA.send', return_value='3EB0') as send:
            self.assertEqual(self.client.post(f'/broadcast/message/{message.id}/send/').json()['status'], 'sent')
        send.assert_called_once_with('77001112233@c.us', message.text)
        # Учителю и бухгалтеру настройки WAHA недоступны
        self.login(self.accountant)
        self.assertEqual(self.client.get('/broadcast/whatsapp/qr/').status_code, 403)


class IncomeStatsTests(CRMTestCase):
    def setUp(self):
        super().setUp()
        from .views import cycles_back
        ScheduleSlot.objects.create(group=self.group, day_of_week=0, start_time=time(10), end_time=time(11))
        Enrollment.objects.update(enrolled_at=timezone.now() - timedelta(days=500))
        year, month = get_current_period()
        # Три месяца назад .. текущий: оплачены все, кроме прошлого месяца
        self.cycles = [y * 100 + m for y, m in cycles_back(year, month, 4)]
        for cycle in (self.cycles[0], self.cycles[1], self.cycles[3]):
            Payment.objects.create(student=self.student, group=self.group, cycle_number=cycle, amount=16000,
                                   paid_amount=16000, is_paid=True, paid_at=timezone.now())
        self.login(self.accountant)

    def test_stats_for_period(self):
        stats = self.client.get('/accountant-stats/?months=6').context['stats']
        self.assertEqual(stats['total'], 48000)
        # График начинается с первого месяца, когда в CRM отметили оплату
        first = stats['chart'][0]
        self.assertEqual(first['year'] * 100 + first['month'], self.cycles[0])
        self.assertTrue(stats['chart'][-1]['is_current'])
        # Долг — только за прошлый месяц; текущий ещё идёт и долгом не считается
        self.assertEqual([d['student'].name for d in stats['debtors']], ['Алия'])
        debt_month = stats['debtors'][0]['months']
        self.assertEqual([m['year'] * 100 + m['month'] for m in debt_month], [self.cycles[2]])
        self.assertGreater(stats['debt'], 0)
        self.assertTrue(0 < stats['rate'] < 100)
        group = next(g for g in stats['groups'] if g['group'] == self.group)
        self.assertEqual((group['collected'], group['debt']), (48000, stats['debt']))
        self.assertEqual(stats['teachers'][0]['name'], 'teacher')

    def test_filters_and_access(self):
        self.assertEqual(self.client.get('/accountant-stats/?months=7').context['months'], 12)
        self.assertEqual(self.client.get('/accountant-stats/?type=individual').context['stats']['total'], 0)
        self.login(self.teacher)
        self.assertRedirects(self.client.get('/accountant-stats/'), '/dashboard/', fetch_redirect_response=False)

    def test_excel_export(self):
        """Excel статистики: те же цифры, что на странице, — формулами из листа «Начисления»"""
        import io
        import openpyxl
        stats = self.client.get('/accountant-stats/?months=6').context['stats']
        response = self.client.get('/accountant-stats/export/?months=6')
        self.assertEqual(response.status_code, 200)
        self.assertIn('enjoy-english-income-6m-', response['Content-Disposition'])
        content = io.BytesIO(response.content)
        formulas, values = openpyxl.load_workbook(content), openpyxl.load_workbook(content, data_only=True)
        self.assertEqual(formulas.sheetnames, ['Сводка', 'По месяцам', 'По группам', 'По учителям', 'Должники', 'Начисления'])
        # Строки начислений: ученик × группа × месяц периода (4 месяца, один ученик)
        charges = values['Начисления']
        last = 4 + sum(1 for r in range(5, 100) if charges.cell(row=r, column=5).value)
        aliya = [charges.cell(row=r, column=9).value for r in range(5, last + 1) if charges.cell(row=r, column=5).value == 'Алия']
        self.assertEqual(aliya, ['Оплачено', 'Оплачено', 'Долг', 'Оплачено'])
        self.assertTrue(formulas['Начисления']['K5'].value.startswith('=IF(I5="Оплачено"'))
        # Сводка — формулы, а рядом готовые числа, совпадающие со страницей
        summary_f, summary_v = formulas['Сводка'], values['Сводка']
        self.assertEqual(summary_f['B7'].value, f"=SUM('Начисления'!$K$5:$K${last})")
        self.assertEqual(summary_v['B7'].value, stats['total'])
        self.assertEqual(summary_v['K7'].value, stats['debt'])               # «Долг»
        self.assertAlmostEqual(summary_v['H7'].value, stats['rate'] / 100, places=2)
        # По месяцам: СУММЕСЛИМН по «Начислениям»
        months = formulas['По месяцам']
        self.assertTrue(months['C5'].value.startswith(f"=SUMIFS('Начисления'!$K$5:$K${last}"))
        self.assertEqual([values['По месяцам'].cell(row=5 + i, column=3).value for i in range(len(stats['chart']))],
                         [c['collected'] for c in stats['chart']])
        self.assertEqual(values['Должники']['A5'].value, 'Алия')
        self.assertEqual(values['Должники']['G5'].value, stats['debt'])
        self.assertEqual(values['По группам']['F5'].value, 48000)
        import zipfile
        charts = [n for n in zipfile.ZipFile(content).namelist() if n.startswith('xl/charts/chart')]
        self.assertEqual(len(charts), 4)                                       # 2 по месяцам, способы оплаты, группы


class CashPaymentTests(CRMTestCase):
    def setUp(self):
        super().setUp()
        ScheduleSlot.objects.create(group=self.group, day_of_week=0, start_time=time(10), end_time=time(11))
        self.year, self.month = get_current_period()
        self.login(self.manager)

    def toggle(self, **extra):
        return self.client.post(f'/payments/group/{self.group.id}/toggle/',
                                {'student_id': self.student.id, 'year': self.year, 'month': self.month, **extra}).json()

    def method(self, method):
        return self.client.post(f'/payments/group/{self.group.id}/method/',
                                {'student_id': self.student.id, 'year': self.year, 'month': self.month, 'method': method}).json()

    def test_mark_cash_and_back(self):
        self.assertFalse(self.method('cash')['success'])          # оплаты ещё нет
        self.assertEqual(self.toggle()['method'], 'bank')         # по умолчанию — банк
        self.assertTrue(self.method('cash')['success'])
        payment = Payment.objects.get()
        self.assertEqual(payment.method, 'cash')
        response = self.client.get(f'/payments/?year={self.year}&month={self.month}')
        self.assertEqual(response.context['cash'], int(payment.paid_amount))
        # Сняли оплату — способ сбрасывается, повторно приняли наличными сразу
        self.toggle()
        self.assertEqual(Payment.objects.get().method, 'bank')
        self.assertEqual(self.toggle(method='cash')['method'], 'cash')
        self.assertEqual(self.client.post(f'/payments/group/{self.group.id}/method/', {'method': 'card'}).status_code, 400)

    def test_cash_share_in_stats(self):
        self.toggle(method='cash')
        self.login(self.admin)
        stats = self.client.get('/accountant-stats/?months=6').context['stats']
        self.assertEqual((stats['cash_share'], stats['bank_share']), (100, 0))
        self.assertEqual(stats['cash'], stats['total'])


class InboxTests(CRMTestCase):
    def send(self, user, student, **extra):
        self.login(user)
        data = {'student_id': student.id, 'category': 'comment', 'text': 'Стал пропускать домашку', **extra}
        return self.client.post('/messages/send/', data, HTTP_X_REQUESTED_WITH='XMLHttpRequest')

    def test_teacher_writes_about_own_student_only(self):
        self.assertTrue(self.send(self.teacher, self.student, category='complaint').json()['success'])
        message = InboxMessage.objects.get()
        self.assertEqual((message.category, message.group, message.author), ('complaint', self.group, self.teacher))
        self.assertEqual(self.send(self.teacher, self.other_student).status_code, 403)
        self.assertFalse(self.send(self.teacher, self.student, text='').json()['success'])
        self.assertEqual(self.send(self.teacher, self.student, category='spam').status_code, 400)
        self.assertEqual(self.send(self.accountant, self.student).status_code, 403)
        self.assertEqual(InboxMessage.objects.count(), 1)
        # Учитель видит отправленное и что оно ещё не прочитано
        self.login(self.teacher)
        sent = self.client.get('/messages/').context
        self.assertFalse(sent['is_inbox'])
        self.assertEqual([m.id for m in sent['sent']], [message.id])

    def test_admin_inbox_counter_and_read(self):
        self.send(self.teacher, self.student)
        self.send(self.manager, self.other_student, category='complaint')
        self.login(self.admin)
        response = self.client.get('/dashboard/')
        self.assertEqual(response.context['inbox_unread'], 2)
        self.assertContains(response, 'class="ee-rail-badge')
        page = self.client.get('/messages/?cat=complaint').context
        self.assertEqual((page['unread_total'], len(page['new'])), (2, 1))
        first = InboxMessage.objects.filter(category='comment').get()
        data = self.client.post(f'/messages/{first.id}/read/').json()
        self.assertEqual((data['read'], data['unread']), (True, 1))
        self.client.post('/messages/read-all/', {'next': 'https://evil.example/'})
        self.assertFalse(InboxMessage.objects.filter(read_at__isnull=True).exists())
        response = self.client.get('/dashboard/')
        self.assertEqual(response.context['inbox_unread'], 0)
        self.assertNotContains(response, 'class="ee-rail-badge')
        # Учителю чужие входящие недоступны
        self.login(self.teacher)
        self.assertEqual(self.client.post(f'/messages/{first.id}/read/').status_code, 403)
        self.assertNotIn('inbox_unread', self.client.get('/dashboard/').context)


    def test_manager_gets_inbox_and_can_write(self):
        self.send(self.teacher, self.student, category='complaint')
        self.login(self.manager)
        response = self.client.get('/dashboard/')
        self.assertEqual(response.context['inbox_unread'], 1)
        self.assertContains(response, 'class="ee-rail-badge')
        page = self.client.get(f'/messages/?student={self.other_student.id}')
        self.assertTrue(page.context['is_inbox'])
        self.assertEqual(len(page.context['new']), 1)
        self.assertContains(page, 'id="composer"')                 # форма у менеджера тоже есть
        self.assertContains(page, '<details class="panel composer compose-fold mb-5" open>')
        message = InboxMessage.objects.get()
        self.assertTrue(self.client.post(f'/messages/{message.id}/read/').json()['read'])
        # Прочитал менеджер — у админа тоже прочитано, и видно кто
        self.login(self.admin)
        page = self.client.get('/messages/')
        self.assertEqual(page.context['unread_total'], 0)
        self.assertContains(page, 'class="mread"')
        self.assertNotContains(page, 'id="composer"')               # админу писать самому себе незачем
        # Менеджер пишет сам — сообщение во входящих
        self.assertTrue(self.send(self.manager, self.other_student).json()['success'])
        self.assertEqual(InboxMessage.objects.filter(author=self.manager).count(), 1)
        # Бухгалтеру вкладка по-прежнему закрыта
        self.login(self.accountant)
        self.assertEqual(self.client.post(f'/messages/{message.id}/read/').status_code, 403)
        self.assertNotIn('inbox_unread', self.client.get('/payments/').context)

class CalibrationMoveTests(CRMTestCase):
    def setUp(self):
        super().setUp()
        self.group.level = 2
        self.group.save()
        self.stronger = Group.objects.create(name='Elementary A2', teacher=self.other_teacher, level=3)
        self.weaker = Group.objects.create(name='Kids Starters', teacher=self.other_teacher, level=1)
        year, month = get_current_period()
        self.cycle = year * 100 + month
        Payment.objects.create(student=self.student, group=self.group, cycle_number=self.cycle, amount=8000,
                               paid_amount=8000, is_paid=True, paid_at=timezone.now())

    def calibrate(self, direction):
        self.login(self.teacher)
        self.client.post('/calibration/create/', {'student_id': self.student.id, 'group_id': self.group.id,
                                                  'direction': direction, 'comment': 'Тянет сильнее'})
        return InboxMessage.objects.get(category='calibration', calibration__direction=direction)

    def test_calibration_suggests_and_moves(self):
        message = self.calibrate('up')
        self.assertEqual(message.text, 'Тянет сильнее')
        self.login(self.admin)
        item = next(m for m in self.client.get('/messages/').context['new'] if m.id == message.id)
        self.assertEqual([g.name for g in item.suggestions], ['Elementary A2'])
        self.client.post(f'/messages/{message.id}/resolve/', {'action': 'move', 'group_id': self.stronger.id})
        self.assertFalse(Enrollment.objects.filter(student=self.student, group=self.group).exists())
        self.assertTrue(Enrollment.objects.filter(student=self.student, group=self.stronger).exists())
        # Оплата за текущий месяц переехала вместе с учеником
        self.assertEqual(Payment.objects.get(student=self.student, cycle_number=self.cycle).group, self.stronger)
        message.refresh_from_db()
        self.assertEqual(message.resolution, 'Переведён в «Elementary A2»')
        self.assertIsNotNone(message.read_at)

    def test_calibration_down_and_keep(self):
        message = self.calibrate('down')
        self.login(self.admin)
        item = next(m for m in self.client.get('/messages/').context['new'] if m.id == message.id)
        self.assertEqual([g.name for g in item.suggestions], ['Kids Starters'])
        self.client.post(f'/messages/{message.id}/resolve/', {'action': 'keep'})
        message.refresh_from_db()
        self.assertEqual(message.resolution, 'Оставлен в группе')
        self.assertTrue(Enrollment.objects.filter(student=self.student, group=self.group).exists())

    def test_guess_level_from_name(self):
        from .views import guess_level
        cases = {'Pre-Intermediate': 4, 'Upper-Intermediate B2': 6, 'Intermediate B1': 5, 'Beginner A1': 2,
                 'Elementary': 3, 'Kids Starters': 1, 'Advanced': 7, 'Speaking club': None}
        self.assertEqual({name: guess_level(name) for name in cases}, cases)


class AuditFixTests(CRMTestCase):
    def setUp(self):
        super().setUp()
        from django.core.cache import cache
        cache.clear()
        self.login(self.admin)

    def test_edit_student_keeps_enrollment_history_and_books(self):
        old = timezone.now() - timedelta(days=300)
        Enrollment.objects.filter(id=self.enrollment.id).update(enrolled_at=old, book_needed=True, book_comment='Нужна Unit 2')
        extra = Group.objects.create(name='Speaking', teacher=self.teacher)
        archived = Group.objects.create(name='Old', teacher=self.teacher, is_active=False)
        Enrollment.objects.create(student=self.student, group=archived)
        self.client.post(f'/student/{self.student.id}/edit/', {
            'name': 'Алия', 'phone': '+7 700 111 22 33', 'price_per_lesson': '4 500', 'groups': [self.group.id, extra.id]})
        kept = Enrollment.objects.get(student=self.student, group=self.group)
        self.assertEqual((kept.enrolled_at, kept.book_needed, kept.book_comment), (old, True, 'Нужна Unit 2'))
        self.assertTrue(Enrollment.objects.filter(student=self.student, group=extra).exists())
        self.assertTrue(Enrollment.objects.filter(student=self.student, group=archived).exists())
        self.student.refresh_from_db()
        self.assertEqual(self.student.price_per_lesson, Decimal('4500'))
        # Сняли галочку — запись убрана
        self.client.post(f'/student/{self.student.id}/edit/', {'name': 'Алия', 'groups': [extra.id]})
        self.assertFalse(Enrollment.objects.filter(student=self.student, group=self.group).exists())

    def test_bad_input_does_not_crash(self):
        response = self.client.post('/students/add/', {'name': 'Новый', 'age': 'abc'})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Student.objects.filter(name='Новый').exists())
        self.assertContains(response, 'value="Новый"')       # введённое не пропало
        self.client.post('/students/add/', {'name': 'Новый', 'price_per_lesson': 'дорого'})
        self.assertFalse(Student.objects.filter(name='Новый').exists())
        self.client.post('/students/add/', {'name': 'Новый', 'groups': ['abc', '999999', str(self.group.id)]})
        self.assertEqual(Student.objects.get(name='Новый').enrollments.get().group, self.group)
        response = self.client.post(f'/student/{self.student.id}/edit/', {'name': '', 'age': '12a'})
        self.assertEqual(response.status_code, 200)
        self.student.refresh_from_db()
        self.assertEqual(self.student.name, 'Алия')
        for method, url, data in [
            ('post', f'/payments/group/{self.group.id}/toggle/', {'student_id': 'abc', 'year': '2026', 'month': '10'}),
            ('post', '/calibration/create/', {'student_id': 'abc', 'group_id': 'abc', 'direction': 'up'}),
            ('post', f'/student/{self.student.id}/add-to-group/', {'groups': ['abc']}),
            ('post', '/groups/add/', {'name': 'G', 'teacher': 'abc'}),
            ('post', f'/group/{self.group.id}/edit/', {'name': 'G', 'teacher': 'abc'}),
            ('get', '/students/?group=abc&age=abc', None),
        ]:
            response = getattr(self.client, method)(url, data) if data else self.client.get(url)
            self.assertLess(response.status_code, 500, url)

    def test_edit_group_validation(self):
        self.client.post(f'/group/{self.group.id}/edit/', {'name': '  ', 'teacher': self.teacher.id})
        self.group.refresh_from_db()
        self.assertEqual(self.group.name, 'Beginner')
        self.client.post(f'/group/{self.group.id}/edit/', {'name': 'Beginner', 'teacher': self.teacher.id,
                                                          'additional_teachers': [self.teacher.id, self.other_teacher.id]})
        self.assertEqual(list(self.group.teachers.all()), [self.other_teacher])

    def test_no_redirect_to_other_sites(self):
        response = self.client.post(f'/student/{self.student.id}/edit/', {'name': 'Алия', 'return_url': 'https://evil.example/'})
        self.assertRedirects(response, '/students/', fetch_redirect_response=False)

    def test_profile_photo_must_be_image(self):
        import io
        from django.core.files.uploadedfile import SimpleUploadedFile
        from PIL import Image
        fake = SimpleUploadedFile('photo.html', b'<script>alert(1)</script>', content_type='image/png')
        self.client.post('/profile/', {'action': 'change_photo', 'photo': fake})
        self.assertFalse(Profile.objects.get(user=self.admin).photo)
        buf = io.BytesIO()
        Image.new('RGB', (8, 8), 'orange').save(buf, 'PNG')
        real = SimpleUploadedFile('me.html', buf.getvalue(), content_type='image/png')
        with self.settings(MEDIA_ROOT=self._tmp_media()):
            self.client.post('/profile/', {'action': 'change_photo', 'photo': real})
            photo = Profile.objects.get(user=self.admin).photo
            self.assertTrue(photo.name.endswith('.png'), photo.name)

    def _tmp_media(self):
        import tempfile
        path = tempfile.mkdtemp()
        self.addCleanup(__import__('shutil').rmtree, path, True)
        return path

    def test_login_throttle(self):
        self.client.logout()
        for _ in range(10):
            self.client.post('/', {'username': 'admin', 'password': 'wrong'})
        response = self.client.post('/', {'username': 'admin', 'password': 'pass'})
        self.assertEqual(response.status_code, 200)               # вход не пустили даже с верным паролем
        self.assertNotIn('_auth_user_id', self.client.session)
        self.assertContains(response, 'Слишком много неудачных попыток')


def mock_attr(obj, name, value):
    from unittest import mock
    return mock.patch.object(obj, name, value)




class QuoteOfTheDayTests(CRMTestCase):
    """Цитата дня на главной (dashboard/quotes.py)"""

    def test_one_quote_per_day(self):
        import re
        from datetime import date
        from .quotes import QUOTES, quote_for
        day = date(2026, 10, 16)
        self.assertEqual(quote_for(day), quote_for(day))
        week = [quote_for(day + timedelta(days=i))['text'] for i in range(len(QUOTES))]
        self.assertEqual(len(set(week)), len(QUOTES))            # пока не прошли все — без повторов
        self.assertEqual(quote_for(day + timedelta(days=len(QUOTES))), quote_for(day))
        self.assertEqual(len({q.ru for q in QUOTES}), len(QUOTES))
        fun = sum(q.kind == 'fun' for q in QUOTES)
        self.assertGreater(fun / len(QUOTES), .6)                # в основном смешные
        self.assertTrue(all(q.kind in ('fun', 'wise') for q in QUOTES))
        # Две мудрые подряд не идут
        kinds = [quote_for(day + timedelta(days=i))['kind'] for i in range(len(QUOTES) + 1)]
        self.assertFalse(any(a == b == 'wise' for a, b in zip(kinds, kinds[1:])))
        cyr = re.compile('[А-Яа-яЁё]')
        for q in QUOTES:
            self.assertTrue(q.ru_author and q.en_author, q.ru)
            self.assertIsNone(cyr.search(q.en + q.en_author), q.en)
        self.assertIsNone(cyr.search(quote_for(day, 'en')['text']))

    def test_dashboard_shows_quote(self):
        from .quotes import quote_for
        quote = quote_for()
        for user in (self.admin, self.teacher, self.accountant, self.manager):
            self.login(user)
            response = self.client.get('/dashboard/')
            self.assertContains(response, 'Цитата дня')
            self.assertContains(response, quote['author'])
        self.client.cookies['ee_lang'] = 'en'
        response = self.client.get('/dashboard/')
        self.assertContains(response, 'Quote of the day')
        self.assertContains(response, quote_for(lang='en')['author'])


class StudentSortTests(CRMTestCase):
    """Ученики: сортировка по дате добавления и подпись даты"""

    def test_plural_and_labels(self):
        from datetime import date, datetime
        from .views import added_info, ru_plural
        self.assertEqual([ru_plural(n, 'день', 'дня', 'дней') for n in (1, 2, 5, 11, 12, 21, 22, 25, 111)],
                         ['день', 'дня', 'дней', 'дней', 'дней', 'день', 'дня', 'дней', 'дней'])
        today = date(2026, 10, 16)
        tz = timezone.get_current_timezone()

        def info(y, m, d, hh=14, mm=30):
            return added_info(datetime(y, m, d, hh, mm, tzinfo=tz), today)
        self.assertEqual(info(2026, 10, 16)['ago'], 'сегодня')
        self.assertEqual(info(2026, 10, 15)['ago'], 'вчера')
        self.assertEqual(info(2026, 10, 13)['ago'], '3 дня назад')
        self.assertEqual(info(2026, 10, 9)['ago'], 'неделю назад')
        self.assertEqual(info(2026, 9, 25)['ago'], '3 недели назад')
        self.assertEqual(info(2026, 9, 16)['ago'], 'месяц назад')
        self.assertEqual(info(2026, 5, 1)['ago'], '5 месяцев назад')
        self.assertEqual(info(2025, 10, 16)['ago'], 'год назад')
        self.assertEqual(info(2023, 1, 1)['ago'], '3 года назад')
        self.assertEqual(info(2026, 10, 16)['label'], '16 окт')
        self.assertEqual(info(2026, 5, 9)['label'], '9 мая')
        self.assertEqual(info(2025, 12, 31)['label'], '31 дек 2025')
        self.assertEqual(info(2026, 10, 16, 9, 5)['full'], 'Добавлен 16.10.2026 в 09:05')
        self.assertTrue(info(2026, 10, 10)['is_new'])
        self.assertFalse(info(2026, 10, 9)['is_new'])

    def test_sorting(self):
        now = timezone.now()
        Student.objects.filter(id=self.student.id).update(created_at=now - timedelta(days=40))       # Алия — старая
        Student.objects.filter(id=self.other_student.id).update(created_at=now - timedelta(days=2))  # Тимур — новый
        Student.objects.create(name='Бахыт')                                                        # только что
        self.login(self.admin)

        def names(response):
            return [d['student'].name for d in response.context['student_data']]
        url = reverse('students_list')
        self.assertEqual(names(self.client.get(url)), ['Алия', 'Бахыт', 'Тимур'])          # по умолчанию — по имени
        response = self.client.get(url + '?sort=new')
        self.assertEqual(names(response), ['Бахыт', 'Тимур', 'Алия'])
        self.assertContains(response, 'class="month-row"')
        self.assertContains(response, 'сегодня')
        self.assertContains(response, '2 дня назад')
        self.assertContains(response, 'class="added-new"')
        # Выбор запоминается
        self.assertEqual(names(self.client.get(url)), ['Бахыт', 'Тимур', 'Алия'])
        self.assertEqual(names(self.client.get(url + '?sort=old')), ['Алия', 'Тимур', 'Бахыт'])
        self.assertEqual(names(self.client.get(url + '?sort=junk')), ['Алия', 'Тимур', 'Бахыт'])
        response = self.client.get(url + '?sort=name')
        self.assertEqual(names(response), ['Алия', 'Бахыт', 'Тимур'])
        self.assertNotContains(response, 'class="month-row"')
        # Сортировка вместе с фильтром: ссылки в заголовках сохраняют фильтр
        response = self.client.get(url + f'?group={self.group.id}&sort=new')
        self.assertEqual(names(response), ['Алия'])
        self.assertIn(f'group={self.group.id}', response.context['sort_urls']['old'])
        # Учитель видит колонку и сортирует своих
        self.login(self.teacher)
        response = self.client.get(url + '?sort=new')
        self.assertContains(response, 'Добавлен')
        self.assertEqual(names(response), ['Алия'])

    def test_english(self):
        Student.objects.filter(id=self.student.id).update(created_at=timezone.now() - timedelta(days=3))
        self.login(self.admin)
        self.client.cookies['ee_lang'] = 'en'
        response = self.client.get(reverse('students_list') + '?sort=new')
        self.assertContains(response, '3 days ago')
        self.assertContains(response, 'Newest first')
        self.assertContains(response, 'Added')


class SetupWhatsAppTests(CRMTestCase):
    """manage.py setup_whatsapp — его вызывает deploy/waha/install.sh, чтобы CRM подключилась к WAHA сама"""

    def run_command(self, *args, env=None):
        import io
        import os
        from unittest import mock
        from django.core.management import call_command
        out = io.StringIO()
        with mock.patch.dict(os.environ, env or {}, clear=False):
            call_command('setup_whatsapp', *args, stdout=out)
        return out.getvalue()

    def test_saves_settings_and_reports_state(self):
        from unittest import mock
        from .models import WhatsAppSettings
        WhatsAppSettings.objects.create(pk=1, provider='green', instance_id='1101', token='old')
        with mock.patch('dashboard.management.commands.setup_whatsapp.waha_status', return_value=('STARTING', '')) as status:
            out = self.run_command('--url', 'http://127.0.0.1:3001/', env={'WAHA_API_KEY': 'k' * 64})
        settings = WhatsAppSettings.load()
        self.assertEqual((settings.provider, settings.api_url, settings.token, settings.instance_id),
                         ('waha', 'http://127.0.0.1:3001', 'k' * 64, ''))
        self.assertTrue(status.called)
        self.assertIn('Осталось отсканировать QR-код', out)
        self.assertNotIn('k' * 64, out)                       # ключ не печатаем
        with mock.patch('dashboard.management.commands.setup_whatsapp.waha_status', return_value=('WORKING', '')):
            self.assertIn('WhatsApp подключён', self.run_command('--status'))

    def test_errors(self):
        import os
        from unittest import mock
        from django.core.management.base import CommandError
        from .whatsapp import WhatsAppError
        with mock.patch.dict(os.environ, {'WAHA_API_KEY': ''}):
            with self.assertRaises(CommandError):
                self.run_command()                            # нет ключа
        with self.assertRaises(CommandError):
            self.run_command('--status')                      # WAHA ещё не настроена
        with mock.patch('dashboard.management.commands.setup_whatsapp.waha_status', side_effect=WhatsAppError('Нет связи')):
            with self.assertRaises(CommandError):
                self.run_command('--key', 'abc')


class SecurityTests(CRMTestCase):
    """dashboard/security.py: IP, подбор пароля, журнал входов, выход, файлы, заголовки, пароли, сеансы"""

    def setUp(self):
        super().setUp()
        from django.core.cache import cache
        cache.clear()

    def test_client_ip_cannot_be_spoofed(self):
        from django.test import RequestFactory
        from .security import client_ip
        rf = RequestFactory()
        outside = rf.get('/', REMOTE_ADDR='8.8.4.4', HTTP_X_FORWARDED_FOR='1.1.1.1', HTTP_X_REAL_IP='2.2.2.2')
        self.assertEqual(client_ip(outside), '8.8.4.4')        # чужим заголовкам не верим
        nginx = rf.get('/', REMOTE_ADDR='127.0.0.1', HTTP_X_REAL_IP='198.51.100.7')
        self.assertEqual(client_ip(nginx), '198.51.100.7')
        chain = rf.get('/', REMOTE_ADDR='127.0.0.1', HTTP_X_FORWARDED_FOR='6.6.6.6, 198.51.100.8')
        self.assertEqual(client_ip(chain), '198.51.100.8')         # последний — дописал наш nginx

    def test_spoofed_header_does_not_bypass_ip_limit(self):
        for i in range(30):
            self.client.post('/', {'username': f'user{i}', 'password': 'x'}, REMOTE_ADDR='8.8.4.4', HTTP_X_FORWARDED_FOR=f'10.0.0.{i}')
        response = self.client.post('/', {'username': 'admin', 'password': 'pass'}, REMOTE_ADDR='8.8.4.4', HTTP_X_FORWARDED_FOR='10.9.9.9')
        self.assertNotIn('_auth_user_id', self.client.session)
        self.assertContains(response, 'Слишком много неудачных попыток')

    def test_login_journal(self):
        from .models import LoginEvent
        self.client.post('/', {'username': 'admin', 'password': 'wrong'}, HTTP_USER_AGENT='Mozilla/5.0 (iPhone) Safari/604')
        self.client.post('/', {'username': 'admin', 'password': 'pass'}, HTTP_USER_AGENT='Mozilla/5.0 (Windows NT 10.0) Chrome/120 Safari/537')
        failed, ok = LoginEvent.objects.order_by('created_at')
        self.assertEqual((failed.success, failed.user, failed.username), (False, None, 'admin'))
        self.assertEqual((ok.success, ok.user), (True, self.admin))
        page = self.client.get('/profile/')
        self.assertContains(page, 'Неверный пароль')
        self.assertContains(page, 'Chrome · Windows')
        self.assertContains(page, 'Safari · iPhone')
        self.assertContains(page, 'Журнал входов в CRM')            # админу — журнал всей CRM
        self.assertContains(page, 'неудачные попытки войти в ваш аккаунт')
        self.login(self.teacher)
        self.assertNotContains(self.client.get('/profile/'), 'Журнал входов в CRM')

    def test_admin_login_is_throttled(self):
        from .models import LoginEvent
        superuser = make_user('boss', superuser=True)
        for _ in range(10):
            self.client.post('/admin/login/', {'username': 'boss', 'password': 'wrong', 'next': '/admin/'})
        response = self.client.post('/admin/login/', {'username': 'boss', 'password': 'pass', 'next': '/admin/'})
        self.assertEqual(response.status_code, 429)
        self.assertNotIn('_auth_user_id', self.client.session)
        self.assertEqual(LoginEvent.objects.filter(username='boss', success=False).count(), 11)
        self.assertTrue(superuser.is_superuser)

    def test_logout_only_by_button(self):
        self.login(self.teacher)
        page = self.client.get('/logout/')                            # ссылкой (или чужой картинкой) — не выходим
        self.assertContains(page, 'Выйти из CRM?')
        self.assertIn('_auth_user_id', self.client.session)
        self.assertRedirects(self.client.post('/logout/'), '/', fetch_redirect_response=False)
        self.assertNotIn('_auth_user_id', self.client.session)
        dashboard = self.client.get('/dashboard/')
        self.assertEqual(dashboard.status_code, 302)

    def test_media_only_for_staff(self):
        import os
        import tempfile
        root = tempfile.mkdtemp()
        self.addCleanup(__import__('shutil').rmtree, root, True)
        for folder in ('logo', 'profile_photos'):
            os.makedirs(os.path.join(root, folder))
            with open(os.path.join(root, folder, 'a.png'), 'wb') as f:
                f.write(b'png')
        with self.settings(MEDIA_ROOT=root):
            self.assertEqual(self.client.get('/media/logo/a.png').status_code, 200)
            self.assertEqual(self.client.get('/media/profile_photos/a.png').status_code, 404)
            self.assertEqual(self.client.get('/media/../config/settings.py').status_code, 404)
            self.login(self.teacher)
            response = self.client.get('/media/profile_photos/a.png')
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response['X-Content-Type-Options'], 'nosniff')

    def test_security_headers_and_secure_cookies(self):
        response = self.client.get('/')
        self.assertIn("object-src 'none'", response['Content-Security-Policy'])
        self.assertIn("frame-ancestors 'self'", response['Content-Security-Policy'])
        self.assertIn('camera=()', response['Permissions-Policy'])
        self.assertNotIn('Strict-Transport-Security', response)
        # За nginx по HTTPS: cookie входа помечаются Secure, браузер запоминает HTTPS
        response = self.client.post('/', {'username': 'admin', 'password': 'pass'},
                                    REMOTE_ADDR='127.0.0.1', HTTP_X_FORWARDED_PROTO='https')
        self.assertTrue(response.cookies['sessionid']['secure'])
        self.assertIn('max-age=', response['Strict-Transport-Security'])
        self.assertTrue(response.cookies['sessionid']['httponly'])
        # Подделанный заголовок снаружи не считается
        self.client.logout()
        response = self.client.post('/', {'username': 'admin', 'password': 'pass'},
                                    REMOTE_ADDR='8.8.4.4', HTTP_X_FORWARDED_PROTO='https')
        self.assertFalse(response.cookies['sessionid']['secure'])

    def test_password_rules(self):
        from .security import password_problems
        self.assertEqual(len(password_problems('12345678')), 2)        # только цифры и слишком простой
        self.assertIn('Пароль — не короче 8 символов', password_problems('Ab3!x'))
        self.assertIn('Пароль слишком похож на логин или имя', password_problems('teacher1', self.teacher))
        self.assertEqual(password_problems('Kitap-Sabak-2026'), [])
        self.login(self.teacher)
        response = self.client.post('/profile/', {'action': 'change_password', 'old_password': 'pass',
                                                   'new_password': 'qwerty', 'confirm_password': 'qwerty'}, follow=True)
        self.assertContains(response, 'Пароль — не короче 8 символов')
        self.teacher.refresh_from_db()
        self.assertTrue(self.teacher.check_password('pass'))
        self.client.post('/profile/', {'action': 'change_password', 'old_password': 'pass',
                                       'new_password': 'Kitap-Sabak-2026', 'confirm_password': 'Kitap-Sabak-2026'})
        self.teacher.refresh_from_db()
        self.assertTrue(self.teacher.check_password('Kitap-Sabak-2026'))

    def test_end_other_sessions(self):
        from django.test import Client
        phone = Client()
        phone.force_login(self.teacher)
        self.login(self.teacher)
        self.assertContains(self.client.get('/profile/'), 'Выйти на других устройствах')
        response = self.client.post('/profile/', {'action': 'end_sessions'}, follow=True)
        self.assertContains(response, 'закрыто сеансов на других устройствах — 1')
        self.assertEqual(phone.get('/dashboard/').status_code, 302)    # на «телефоне» вышли
        self.assertEqual(self.client.get('/dashboard/').status_code, 200)

    def test_device_label(self):
        from .security import device_label
        self.assertEqual(device_label('Mozilla/5.0 (Linux; Android 13) AppleWebKit Chrome/120 Mobile Safari'), 'Chrome · Android')
        self.assertEqual(device_label('Mozilla/5.0 (Windows NT 10.0) Chrome/120 Safari/537 Edg/120'), 'Edge · Windows')
        self.assertEqual(device_label(''), 'Неизвестное устройство')


class GeneralExcelTests(CRMTestCase):
    """Общий отчёт в Excel (/export/): ученики, оплаты, группы, журнал посещаемости — формулами"""

    def test_general_report(self):
        import io
        import openpyxl
        ScheduleSlot.objects.create(group=self.group, day_of_week=0, start_time=time(10), end_time=time(11))
        lesson = Lesson.objects.create(group=self.group, date=timezone.localdate())
        Attendance.objects.create(lesson=lesson, student=self.student, status='present', homework_done=True, stars=3)
        cancelled = Lesson.objects.create(group=self.group, date=timezone.localdate() - timedelta(days=1), is_cancelled=True)
        Attendance.objects.create(lesson=cancelled, student=self.student, status='absent')
        year, month = get_current_period(self.group.cycle_start_day)
        Payment.objects.create(student=self.student, group=self.group, cycle_number=year * 100 + month,
                               amount=8000, paid_amount=8000, is_paid=True)
        self.login(self.manager)
        response = self.client.get('/export/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('enjoy-english-report-', response['Content-Disposition'])
        content = io.BytesIO(response.content)
        formulas, values = openpyxl.load_workbook(content), openpyxl.load_workbook(content, data_only=True)
        self.assertEqual(formulas.sheetnames, ['Сводка', 'Ученики', 'Оплаты', 'Группы', 'Посещаемость'])
        # Журнал: только проведённые уроки
        log = values['Посещаемость']
        self.assertEqual((log['E5'].value, log['F5'].value, log['G5'].value, log['H5'].value), ('Алия', 'Был', 'Да', 3))
        self.assertIsNone(log['E6'].value)
        # Ученики: посещаемость и звёзды — формулами из журнала
        students = {values['Ученики'].cell(row=r, column=1).value: r for r in range(5, 7)}
        r = students['Алия']
        self.assertTrue(formulas['Ученики'][f'N{r}'].value.startswith('=IFERROR(M'))
        self.assertEqual((values['Ученики'][f'N{r}'].value, values['Ученики'][f'P{r}'].value), (1, 3))
        # Оплаты: текущий месяц оплачен, к оплате = цена × уроки
        pay = {values['Оплаты'].cell(row=r, column=1).value: r for r in range(5, 7)}
        r = pay['Алия']
        self.assertEqual(formulas['Оплаты'][f'F{r}'].value, f'=D{r}*E{r}')
        self.assertEqual(values['Оплаты'][f'G{r}'].value, 'Оплачено')
        # Группы и сводка
        groups = {values['Группы'].cell(row=r, column=1).value: r for r in range(5, 7)}
        self.assertEqual(values['Группы'][f'G{groups["Beginner"]}'].value, 1)
        self.assertEqual(values['Сводка']['B6'].value, 2)                       # учатся сейчас
        self.assertTrue(formulas['Сводка']['B6'].value.startswith('=COUNTIFS('))
        import zipfile
        self.assertGreaterEqual(len([n for n in zipfile.ZipFile(content).namelist() if n.startswith('xl/charts/chart')]), 3)
        # Учителю — нельзя
        self.login(self.teacher)
        self.assertEqual(self.client.get('/export/').status_code, 302)

    def test_empty_crm_exports(self):
        """Ни учеников, ни оплат — файлы всё равно собираются"""
        Payment.objects.all().delete()
        Student.objects.all().delete()
        Group.objects.all().delete()
        self.login(self.admin)
        self.assertEqual(self.client.get('/export/').status_code, 200)
        self.assertEqual(self.client.get('/accountant-stats/export/?months=12&type=individual').status_code, 200)

class BackupTests(TransactionTestCase):
    """Без обёртки-транзакции TestCase: SQLite не даёт снять копию посреди открытой записи"""
    def setUp(self):
        self.admin = make_user('admin', 'admin')
        self.teacher = make_user('teacher', 'teacher')

    def test_backup_command_and_download(self):
        import sqlite3
        import tempfile
        from pathlib import Path
        from django.core.management import call_command
        from dashboard.management.commands import backup_db
        target = Path(tempfile.mkdtemp())
        self.addCleanup(__import__('shutil').rmtree, target, True)
        with mock_attr(backup_db, 'BACKUP_DIR', target):
            call_command('backup_db', '--keep', '1', stdout=open('/dev/null', 'w'))
            call_command('backup_db', '--keep', '1', stdout=open('/dev/null', 'w'))
        files = list(target.glob('db-*.sqlite3'))
        self.assertEqual(len(files), 1)
        sqlite3.connect(str(files[0])).execute('select count(*) from sqlite_master').fetchone()
        self.client.force_login(self.admin)
        response = self.client.get('/backup/download/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertEqual(b''.join(response.streaming_content)[:15], b'SQLite format 3')
        self.client.force_login(self.teacher)
        self.assertRedirects(self.client.get('/backup/download/'), '/dashboard/', fetch_redirect_response=False)

